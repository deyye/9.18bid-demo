from fastapi import APIRouter, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
from ..models.schemas import FileUploadResponse, AnalysisRequest, AnalysisType, AnalysisResponse, ExportRequest, OutlineItem
from ..services.file_service import FileService
from ..services.openai_service import OpenAIService
from ..services.qwen_api import QwenService
from ..utils.config_manager import config_manager
import json
import re
from io import BytesIO

# 补充 python-docx 和 io 的引用
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from typing import List

router = APIRouter(prefix="/api/document", tags=["文档处理"])


@router.post("/upload", response_model=FileUploadResponse)
async def upload_file(file: UploadFile = File(...)):
    """上传文档文件并提取文本内容"""
    try:
        # 检查文件类型
        allowed_types = [
            "application/pdf",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ]
        
        if file.content_type not in allowed_types:
            return FileUploadResponse(
                success=False,
                message="不支持的文件类型，请上传PDF或Word文档"
            )
        
        # 处理文件并提取文本
        file_content = await FileService.process_uploaded_file(file)
        
        return FileUploadResponse(
            success=True,
            message=f"文件 {file.filename} 上传成功",
            file_content=file_content
        )
        
    except Exception as e:
        return FileUploadResponse(
            success=False,
            message=f"文件处理失败: {str(e)}"
        )

# ------------------------------
# 改造文档分析接口：支持Qwen/OpenAI切换
# ------------------------------
@router.post("/analyze-stream")
async def analyze_document(
    request: AnalysisRequest,
    use_qwen: bool = True,
    stream: bool = True  # 新增参数，是否启用流式
):
    """分析文档内容，支持流式/非流式输出"""
    try:
        # 1. 初始化模型服务
        if use_qwen:
            model_service = QwenService(base_url="http://localhost:10086")
        else:
            config = config_manager.load_config()
            if not config.get("api_key"):
                raise HTTPException(status_code=400, detail="请先配置OpenAI API密钥")
            model_service = OpenAIService(
                api_key=config["api_key"],
                base_url=config.get("base_url", ""),
                model_name=config.get("model_name", "gpt-3.5-turbo")
            )

        # 2. 构建分析提示词
        if request.analysis_type == AnalysisType.OVERVIEW:
            system_prompt = """你是一个专业的标书撰写专家。请按照以下结构化框架分析用户提供的招标文件，提取并总结项目概述信息：
            1. 项目基础信息提取
                - 项目全称（含招标编号）
                - 招标单位名称及地址
                - 项目实施地点（具体到区县/园区/建筑编号）
                - 项目所属行业领域（如智慧交通/医疗信息化等）
            2. 项目背景分析
                - 政策依据（注明具体文件名称及文号）
                - 建设必要性说明（直接引用招标文件中"项目背景"章节原文）
                - 预期社会效益（需包含量化指标如服务人群数量、效率提升百分比等）
            3. 规模参数采集
                - 建设规模（明确建筑面积/设备数量/服务范围等具体数值）
                - 投资总额（需注明金额、货币单位及是否含税）
                - 分项预算构成（要求列出硬件、软件、实施、运维等分项金额）
            4. 时间节点解析
                - 项目启动日期（精确到年月日）
                - 关键里程碑（含设计、开发、测试、验收等阶段的截止日期）
                - 最终交付期限（需注明具体时点及逾期处罚条款）
            5. 实施内容拆解
                - 系统架构图（要求提取并转述技术拓扑结构）
                - 功能模块清单（需包含至少5个核心功能模块的名称及技术参数）
                - 工程量清单（要求列出主要设备型号、数量、安装位置）
            6. 技术特征识别
                - 技术标准要求（需注明遵循的国家标准/行业标准编号）
                - 关键技术指标（要求列出响应时间、并发处理量等量化参数）
                - 创新技术应用（需区分现有技术与拟采用的专利技术）
            7. 额外约束条件
                - 质量验收标准（需引用具体检测规范编号）
                - 安全合规要求（要求列出网络安全等级、数据加密标准等）
                - 运维服务条款（需注明服务响应时间、故障处理时限等）
            工作规范：
                1. 信息提取需满足以下标准：
                    - 每个维度提取不少于3个关键数据点
                    - 技术参数需保留原始单位和精度要求
                    - 时间节点需转换为标准日期格式（YYYY-MM-DD）
                2. 严格排除以下商务要素：
                    - 合同金额及支付条款
                    - 投标资质要求
                    - 评标办法及评分细则
                    - 法律责任条款
                3. 输出格式要求：
                    - 采用分级标题结构（H1-H3）
                    - 每项内容使用项目符号列表
                    - 重要参数使用【】标注
                    - 禁止使用任何Markdown格式
                    - 严格遵循"提取-转述-标注"的三段式结构
"""
        else:  # 技术评分要求提取
            system_prompt = """你是高级技术战略顾问和标书撰写专家，擅长从复杂的招标文档中高效提取“技术评分项”相关内容。请严格按照以下步骤和规则执行任务：
            ### 1. 目标定位
                - 明确识别并提取文档中包含“技术评分”、“评标方法”、“评分标准”、“技术参数”、“技术要求”、“技术方案”、“技术部分”或“评审要素”的章节（如“第5章 评标方法”或“附件3-技术评分表”），需完整列出章节编号及标题。
                - 排除所有与商务条款、价格报价、企业资质、服务承诺等非技术类评分项相关的段落或表格。
            ### 2. 提取内容要求
                对每一项技术评分项，按以下结构化格式输出（若信息缺失，标注“未提及”），若评分项描述模糊需结合上下文进行逻辑推导并补充完整：
                【评分项名称】：<原文中精确的技术评分项名称，保留专业术语及原文表述>
                【权重/分值】：<具体分值或占比数值，需统一为“分”或“%”单位，并标注原文单位（如“[原文：20点]”或“[原文：15%]”）>
                【评分标准】：<完整描述评分规则，包括合格线、扣分/加分细则、最高/最低分限制等，若原文未明确则标注“未提及”>
                【数据来源】：<文档中精确的位置信息，需包含章节编号、附件编号、页码或条款编号（如“第5.2.3条”或“附件3-表2（第3页）”）>
            ### 3. 处理规则
                - 模糊表述处理：当文档未明确标注“技术评分表”时，需通过分析“技术方案”、“技术要求”等章节中包含“得分”、“评分”、“评审”等关键词的段落，推导出隐含的技术评分项并标注“[推导项]”。
                - 表格数据提取：若评分项以表格形式呈现，需按表格行逐项提取，并在数据来源标注“[表格数据]”标识，同时保留表格中所有技术参数列。
                - 分层结构处理：若存在二级评分项（如“技术方案→子项1、子项2”），需使用“→”符号明确层级关系，并在数据来源标注对应层级编号（如“第6.1.2条→第6.1.2.1款”）。
                - 单位统一规范：将所有分值转换为“分”或“%”单位，若原文使用非标准单位（如“20点”），需在括号中标注原文单位并注明转换依据（如“[原文：20点]（按1点=1分转换）”）。
            ### 4. 输出示例
                【评分项名称】：系统可用性 
                【权重/分值】：25分 [原文：25分] 
                【评分标准】：年平均故障时间≤1小时得满分；每增加1小时扣2分，最高扣10分。 
                【数据来源】：附件4-技术评分细则（第3页） 
                【评分项名称】：响应时间 
                【权重/分值】：15分 [原文：15%] 
                【评分标准】：≤50ms得满分；每增加10ms扣1分。 
                【数据来源】：第6.1.2条→第6.1.2.1款 
            ### 5. 验证步骤
                提取完成后，执行以下强制性自检：
                - [ ] 核对文档中所有包含“技术评分”、“评审要素”等关键词的段落是否被完整覆盖，确保无遗漏章节或条款。
                - [ ] 验证所有提取项的权重总和是否与文档声明的技术分总分一致（如“技术部分共60分”），若存在矛盾需标注“[权重冲突：原文声明XX分，提取总分XX分]”并说明差异。
"""

        # 3. 构造对话列表
        analysis_type_cn = "项目概述" if request.analysis_type == AnalysisType.OVERVIEW else "技术评分要求"
        user_prompt = f"请分析以下招标文件内容，提取{analysis_type_cn}信息：\n\n{request.file_content}"

        messages = [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ]

        # 3. 根据 stream 参数选择调用方式
        if stream:
            async def event_generator():
                full_content = ""  # 用于累积所有片段，生成完整结果
                async for chunk in model_service.chat_completion_stream(messages, temperature=0.3):
                    
                    # 实时推送当前片段
                    yield f"data: {json.dumps({'chunk': chunk}, ensure_ascii=False)}\n\n"
                    
                    # 累积片段到完整内容
                    full_content += chunk
                
                # 所有片段生成完毕后，推送完整结果
                if full_content:  # 确保有内容才推送
                    yield f"data: {json.dumps({'full': full_content}, ensure_ascii=False)}\n\n"
                
                # 发送结束信号
                # yield "data: [DONE]\n\n"
            
            return StreamingResponse(
                event_generator(),
                media_type="text/event-stream",
                headers={
                    "Cache-Control": "no-cache",
                    "Connection": "keep-alive",
                }
            )

        else:
            # 非流式调用
            full_result = await model_service.chat_completion(messages, temperature=0.3)
            return AnalysisResponse(
                success=True,
                message="文档分析完成",
                result=full_result
            )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"文档分析失败: {str(e)}")
    
@router.post("/export")
async def export_document(request: ExportRequest):
    """
    将生成的内容和目录导出为 Word 文档
    """
    try:
        # 创建 Word 文档
        doc = Document()
        
        # 设置中文字体辅助函数
        def set_font(run, font_name='宋体', size=None):
            run.font.name = font_name
            run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
            if size:
                run.font.size = Pt(size)

        # 添加标题
        title = doc.add_heading(level=0)
        run = title.add_run("投标文件")
        set_font(run, '黑体', 24)
        title.alignment = 1 # 居中

        doc.add_page_break()

        # 递归添加章节内容
        def add_chapter(items: List[OutlineItem], level: int = 1):
            for item in items:
                # 添加章节标题
                # Word 标题等级最多到 9
                heading_level = level if level <= 9 else 9
                heading = doc.add_heading(level=heading_level)
                run = heading.add_run(item.title)
                # 根据层级简单设置字体
                if level == 1:
                    set_font(run, '黑体', 16)
                else:
                    set_font(run, '黑体', 14)

                # 获取并清理内容
                content = request.content.get(item.id, "")
                if content:
                    # 简单去除 HTML 标签（Quill 返回的是 HTML）
                    # 替换常见块级标签为换行
                    text = content.replace("</p>", "\n").replace("<p>", "")
                    text = text.replace("<br>", "\n").replace("</h1>", "\n").replace("</h2>", "\n")
                    # 去除所有其他 HTML 标签
                    text = re.sub(r'<[^>]+>', '', text).strip()
                    
                    # 添加段落
                    if text:
                        # 处理多段落
                        for para_text in text.split('\n'):
                            if para_text.strip():
                                p = doc.add_paragraph(para_text.strip())
                                p.paragraph_format.first_line_indent = Pt(24) # 首行缩进
                                set_font(p.add_run(para_text.strip()), '宋体', 12)

                # 递归处理子章节
                if item.children:
                    add_chapter(item.children, level + 1)

        add_chapter(request.outline)

        # 保存到内存流
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        # 返回流式响应
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={
                "Content-Disposition": "attachment; filename=bid_document.docx",
                "Access-Control-Expose-Headers": "Content-Disposition"
            }
        )

    except Exception as e:
        print(f"导出失败: {e}")
        raise HTTPException(status_code=500, detail=f"导出失败: {str(e)}")